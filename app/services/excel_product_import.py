from __future__ import annotations

import io
import json
import uuid
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import Select, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.category import Category
from app.models.category_characteristic import CategoryCharacteristic
from app.models.characteristic import CharacteristicType, GlobalCharacteristic

STEP1_HEADERS = [
    "Название товара",
    "Категория",
]

STEP2_BASE_HEADERS = [
    "import_row_id",
    "Название товара",
    "Категория",
    "Цена",
    "Описание",
    "Объём, м³",
    "Вес, кг",
    "Скидка, %",
]

META_VERSION = "v1"


@dataclass
class ImportErrorItem:
    sheet: str
    row: int
    column: str
    message: str


@dataclass
class Step1Row:
    row_number: int
    import_row_id: str
    category_id: uuid.UUID
    category_path: str
    name: str
    price: int
    description: str | None
    volume_m3: float
    weight_kg: float
    is_active: bool
    discount_percent: int | None


@dataclass
class Step2Row:
    sheet_name: str
    row_number: int
    import_row_id: str
    base_row: Step1Row
    characteristic_values: dict[int, tuple[str | None, float | None]]


@dataclass
class CharacteristicMeta:
    characteristic_id: int
    value_type: str
    column_index: int


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_category_token(token: str) -> str:
    return " ".join(token.strip().split()).lower()


async def _load_all_categories(session: AsyncSession) -> list[Category]:
    res = await session.execute(select(Category).order_by(Category.name.asc()))
    return list(res.scalars().all())


def _build_tree_maps(categories: list[Category]) -> tuple[dict[uuid.UUID | None, list[Category]], dict[uuid.UUID, Category]]:
    by_parent: dict[uuid.UUID | None, list[Category]] = {}
    by_id: dict[uuid.UUID, Category] = {}
    for cat in categories:
        by_id[cat.id] = cat
        by_parent.setdefault(cat.parent_id, []).append(cat)
    return by_parent, by_id


def _build_category_path(category_id: uuid.UUID, by_id: dict[uuid.UUID, Category]) -> str:
    names: list[str] = []
    current = by_id.get(category_id)
    while current:
        names.append(current.name)
        current = by_id.get(current.parent_id) if current.parent_id else None
    return " / ".join(reversed(names))


def _resolve_category_by_path(path_value: str, by_parent: dict[uuid.UUID | None, list[Category]]) -> Category | None:
    tokens = [_normalize_category_token(t) for t in path_value.split("/") if _normalize_category_token(t)]
    if not tokens:
        return None

    parent_id: uuid.UUID | None = None
    current: Category | None = None

    for token in tokens:
        options = by_parent.get(parent_id, [])
        matched = None
        for option in options:
            if _normalize_category_token(option.name) == token:
                matched = option
                break
        if not matched:
            return None
        current = matched
        parent_id = matched.id

    return current


def _is_leaf(category: Category, by_parent: dict[uuid.UUID | None, list[Category]]) -> bool:
    return len(by_parent.get(category.id, [])) == 0


def _to_int(raw: Any) -> int:
    if raw is None or str(raw).strip() == "":
        raise ValueError("пустое значение")
    if isinstance(raw, bool):
        raise ValueError("ожидалось число")
    try:
        value = int(Decimal(str(raw).strip().replace(",", ".")))
    except (InvalidOperation, ValueError):
        raise ValueError("ожидалось целое число")
    if value < 0:
        raise ValueError("должно быть >= 0")
    return value


def _to_float(raw: Any) -> float:
    if raw is None or str(raw).strip() == "":
        raise ValueError("пустое значение")
    try:
        value = float(str(raw).strip().replace(",", "."))
    except ValueError:
        raise ValueError("ожидалось число")
    if value < 0:
        raise ValueError("должно быть >= 0")
    return value


def _to_bool(raw: Any, default: bool = True) -> bool:
    text = _normalize_text(raw).lower()
    if not text:
        return default
    if text in {"1", "true", "yes", "да", "активен"}:
        return True
    if text in {"0", "false", "no", "нет", "не активен"}:
        return False
    raise ValueError("ожидалось да/нет")


def _to_optional_int(raw: Any) -> int | None:
    text = _normalize_text(raw)
    if not text:
        return None
    val = _to_int(text)
    if val > 100:
        raise ValueError("должно быть <= 100")
    return val


def _set_headers(ws, headers: list[str]) -> None:
    ws.append(headers)
    for idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(idx)].width = max(18, len(title) + 2)


async def generate_step1_template(session: AsyncSession) -> bytes:
    categories = await _load_all_categories(session)
    by_parent, by_id = _build_tree_maps(categories)

    leaf_paths = [
        _build_category_path(cat.id, by_id)
        for cat in categories
        if _is_leaf(cat, by_parent)
    ]
    leaf_paths.sort()

    wb = Workbook()
    ws_products = wb.active
    ws_products.title = "Товары"
    _set_headers(ws_products, STEP1_HEADERS)

    ws_refs = wb.create_sheet("Справочники")
    ws_refs["A1"] = "category_path"
    ws_refs["A1"].font = Font(bold=True)
    for idx, path in enumerate(leaf_paths, start=2):
        ws_refs.cell(row=idx, column=1, value=path)

    last = max(2, len(leaf_paths) + 1)
    dv = DataValidation(type="list", formula1=f"=Справочники!$A$2:$A${last}", allow_blank=False)
    ws_products.add_data_validation(dv)
    dv.add("B2:B1048576")

    ws_refs.sheet_state = "hidden"

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


async def parse_step1_build_step2(session: AsyncSession, source_bytes: bytes) -> tuple[bytes | None, list[ImportErrorItem]]:
    errors: list[ImportErrorItem] = []

    try:
        wb = load_workbook(io.BytesIO(source_bytes), data_only=True)
    except Exception:
        return None, [ImportErrorItem(sheet="Товары", row=0, column="file", message="Не удалось прочитать Excel-файл.")]

    if "Товары" not in wb.sheetnames:
        return None, [ImportErrorItem(sheet="Товары", row=0, column="sheet", message="Не найден лист 'Товары'.")]

    ws = wb["Товары"]
    categories = await _load_all_categories(session)
    by_parent, by_id = _build_tree_maps(categories)

    valid_rows: list[Step1Row] = []

    for row_idx in range(2, ws.max_row + 1):
        row_values = [ws.cell(row=row_idx, column=i).value for i in range(1, len(STEP1_HEADERS) + 1)]
        if not any(_normalize_text(v) for v in row_values):
            continue

        name = _normalize_text(row_values[0])
        category_path_raw = _normalize_text(row_values[1])

        if not name:
            errors.append(ImportErrorItem("Товары", row_idx, "Название товара", "Поле обязательно."))
        if not category_path_raw:
            errors.append(ImportErrorItem("Товары", row_idx, "Категория", "Поле обязательно."))

        category = _resolve_category_by_path(category_path_raw, by_parent) if category_path_raw else None
        if category_path_raw and not category:
            errors.append(ImportErrorItem("Товары", row_idx, "Категория", "Категория по пути не найдена."))
        if category and not _is_leaf(category, by_parent):
            errors.append(ImportErrorItem("Товары", row_idx, "Категория", "Можно указывать только листовую категорию."))

        if not errors or all(er.row != row_idx for er in errors):
            assert category is not None
            valid_rows.append(
                Step1Row(
                    row_number=row_idx,
                    import_row_id=str(uuid.uuid4()),
                    category_id=category.id,
                    category_path=_build_category_path(category.id, by_id),
                    name=name,
                    price=0,
                    description=None,
                    volume_m3=0.0,
                    weight_kg=0.0,
                    is_active=False,
                    discount_percent=None,
                )
            )

    if errors:
        return None, errors

    step2_bytes = await _build_step2_workbook(session, valid_rows, by_id)
    return step2_bytes, []


async def _load_characteristics_for_categories(session: AsyncSession, category_ids: list[uuid.UUID]) -> dict[uuid.UUID, list[GlobalCharacteristic]]:
    if not category_ids:
        return {}
    stmt: Select = (
        select(CategoryCharacteristic.category_id, GlobalCharacteristic)
        .join(GlobalCharacteristic, GlobalCharacteristic.id == CategoryCharacteristic.characteristic_id)
        .where(CategoryCharacteristic.category_id.in_(category_ids))
        .order_by(CategoryCharacteristic.category_id, GlobalCharacteristic.name.asc())
    )
    res = await session.execute(stmt)
    grouped: dict[uuid.UUID, list[GlobalCharacteristic]] = {}
    for cat_id, ch in res.all():
        grouped.setdefault(cat_id, []).append(ch)
    return grouped


def _safe_sheet_name(raw: str, used: set[str]) -> str:
    bad = set("[]:*?/\\")
    cleaned = "".join(ch for ch in raw if ch not in bad).strip() or "Категория"
    cleaned = cleaned[:31]
    candidate = cleaned
    i = 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        i += 1
    used.add(candidate)
    return candidate


async def _build_step2_workbook(session: AsyncSession, rows: list[Step1Row], by_id: dict[uuid.UUID, Category]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    rows_by_cat: dict[uuid.UUID, list[Step1Row]] = {}
    for row in rows:
        rows_by_cat.setdefault(row.category_id, []).append(row)

    ch_by_cat = await _load_characteristics_for_categories(session, list(rows_by_cat.keys()))

    meta_ws = wb.create_sheet("META")
    meta_ws.sheet_state = "hidden"
    meta_ws["A1"] = "version"
    meta_ws["B1"] = META_VERSION
    meta_ws["A3"] = "column_mapping"
    meta_ws.append(["sheet_name", "column_index", "characteristic_id", "characteristic_value_type"])

    used_names: set[str] = {"META"}
    sheet_name_by_category: dict[uuid.UUID, str] = {}

    for category_id, cat_rows in rows_by_cat.items():
        category_path = _build_category_path(category_id, by_id)
        sheet_name = _safe_sheet_name(category_path, used_names)
        sheet_name_by_category[category_id] = sheet_name

        ws = wb.create_sheet(sheet_name)
        base_headers = list(STEP2_BASE_HEADERS)
        cat_characteristics = ch_by_cat.get(category_id, [])
        ch_headers = [f"{ch.name} ({ch.unit})" if ch.unit else ch.name for ch in cat_characteristics]
        all_headers = base_headers + ch_headers
        _set_headers(ws, all_headers)

        for base_col_idx in range(1, len(base_headers) + 1):
            if base_col_idx == 1:
                ws.column_dimensions[get_column_letter(base_col_idx)].hidden = True

        for item in cat_rows:
            ws.append(
                [
                    item.import_row_id,
                    item.name,
                    item.category_path,
                    None,
                    None,
                    None,
                    None,
                    None,
                    *([None] * len(cat_characteristics)),
                ]
            )

        first_ch_col = len(base_headers) + 1
        for offset, ch in enumerate(cat_characteristics):
            col_idx = first_ch_col + offset
            meta_ws.append([sheet_name, col_idx, ch.id, ch.value_type.value])

    meta_ws.append([])
    meta_ws.append(["row_mapping"])
    meta_ws.append(
        [
            "import_row_id",
            "sheet_name",
            "category_id",
            "category_path",
            "name",
        ]
    )
    for item in rows:
        meta_ws.append(
            [
                item.import_row_id,
                sheet_name_by_category[item.category_id],
                str(item.category_id),
                item.category_path,
                item.name,
            ]
        )

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def parse_step2(source_bytes: bytes) -> tuple[dict[str, list[CharacteristicMeta]], dict[str, Step1Row], list[Step2Row], list[ImportErrorItem]]:
    errors: list[ImportErrorItem] = []
    try:
        wb = load_workbook(io.BytesIO(source_bytes), data_only=True)
    except Exception:
        return {}, {}, [], [ImportErrorItem("META", 0, "file", "Не удалось прочитать Excel-файл.")]

    if "META" not in wb.sheetnames:
        return {}, {}, [], [ImportErrorItem("META", 0, "sheet", "Не найден служебный лист META.")]

    meta = wb["META"]
    if _normalize_text(meta["B1"].value) != META_VERSION:
        return {}, {}, [], [ImportErrorItem("META", 1, "version", "Неверная версия файла импорта.")]

    mapping_by_sheet: dict[str, list[CharacteristicMeta]] = {}
    row_map: dict[str, Step1Row] = {}

    mode = None
    for r in range(1, meta.max_row + 1):
        a = _normalize_text(meta.cell(r, 1).value)
        if a == "column_mapping":
            mode = "column"
            continue
        if a == "row_mapping":
            mode = "row"
            continue
        if a in {"sheet_name", "import_row_id", ""}:
            continue

        if mode == "column":
            sheet_name = _normalize_text(meta.cell(r, 1).value)
            if not sheet_name:
                continue
            try:
                col_idx = int(meta.cell(r, 2).value)
                characteristic_id = int(meta.cell(r, 3).value)
            except Exception:
                errors.append(ImportErrorItem("META", r, "mapping", "Некорректный mapping характеристик."))
                continue
            vtype = _normalize_text(meta.cell(r, 4).value) or "string"
            mapping_by_sheet.setdefault(sheet_name, []).append(
                CharacteristicMeta(characteristic_id=characteristic_id, value_type=vtype, column_index=col_idx)
            )

        elif mode == "row":
            import_row_id = _normalize_text(meta.cell(r, 1).value)
            if not import_row_id:
                continue
            try:
                category_id = uuid.UUID(_normalize_text(meta.cell(r, 3).value))
            except Exception:
                errors.append(ImportErrorItem("META", r, "category_id", "Некорректный category_id в row mapping."))
                continue

            row_map[import_row_id] = Step1Row(
                row_number=0,
                import_row_id=import_row_id,
                category_id=category_id,
                category_path=_normalize_text(meta.cell(r, 4).value),
                name=_normalize_text(meta.cell(r, 5).value),
                price=0,
                description=None,
                volume_m3=0.0,
                weight_kg=0.0,
                is_active=False,
                discount_percent=None,
            )

    parsed_rows: list[Step2Row] = []
    for sheet_name in wb.sheetnames:
        if sheet_name == "META":
            continue
        ws = wb[sheet_name]
        mapping = mapping_by_sheet.get(sheet_name, [])

        for r in range(2, ws.max_row + 1):
            import_row_id = _normalize_text(ws.cell(r, 1).value)
            if not import_row_id:
                continue
            if import_row_id not in row_map:
                errors.append(ImportErrorItem(sheet_name, r, "import_row_id", "Не найден import_row_id в META."))
                continue

            name = _normalize_text(ws.cell(r, 2).value)
            if not name:
                errors.append(ImportErrorItem(sheet_name, r, "Название товара", "Поле обязательно."))

            try:
                price = _to_int(ws.cell(r, 4).value)
            except ValueError as e:
                errors.append(ImportErrorItem(sheet_name, r, "Цена", str(e)))
                price = 0

            description = _normalize_text(ws.cell(r, 5).value) or None

            try:
                volume_m3 = _to_float(ws.cell(r, 6).value)
            except ValueError as e:
                errors.append(ImportErrorItem(sheet_name, r, "Объём, м³", str(e)))
                volume_m3 = 0.0

            try:
                weight_kg = _to_float(ws.cell(r, 7).value)
            except ValueError as e:
                errors.append(ImportErrorItem(sheet_name, r, "Вес, кг", str(e)))
                weight_kg = 0.0

            try:
                discount_percent = _to_optional_int(ws.cell(r, 8).value)
            except ValueError as e:
                errors.append(ImportErrorItem(sheet_name, r, "Скидка, %", str(e)))
                discount_percent = None

            meta_row = row_map[import_row_id]
            base_row = Step1Row(
                row_number=r,
                import_row_id=import_row_id,
                category_id=meta_row.category_id,
                category_path=meta_row.category_path,
                name=name,
                price=price,
                description=description,
                volume_m3=volume_m3,
                weight_kg=weight_kg,
                is_active=False,
                discount_percent=discount_percent,
            )

            ch_values: dict[int, tuple[str | None, float | None]] = {}
            for ch in mapping:
                raw = ws.cell(r, ch.column_index).value
                txt = _normalize_text(raw)
                if not txt:
                    continue

                if ch.value_type == CharacteristicType.number.value:
                    try:
                        num = _to_float(raw)
                    except ValueError:
                        errors.append(
                            ImportErrorItem(
                                sheet_name,
                                r,
                                ws.cell(1, ch.column_index).value or str(ch.column_index),
                                "ожидалось число",
                            )
                        )
                        continue
                    ch_values[ch.characteristic_id] = (None, num)
                else:
                    ch_values[ch.characteristic_id] = (txt, None)

            parsed_rows.append(
                Step2Row(
                    sheet_name=sheet_name,
                    row_number=r,
                    import_row_id=import_row_id,
                    base_row=base_row,
                    characteristic_values=ch_values,
                )
            )

    return mapping_by_sheet, row_map, parsed_rows, errors


def dump_errors_as_text(errors: list[ImportErrorItem]) -> str:
    return json.dumps([e.__dict__ for e in errors], ensure_ascii=False, indent=2)
